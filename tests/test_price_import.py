"""Toplu fiyat (Excel): the supplier's workbook read block by block, matched by code, written to all four groups."""

import io
import uuid
import zipfile
from decimal import Decimal

import openpyxl
import pytest
from httpx import AsyncClient
from sqlalchemy import delete, select
from sqlalchemy.ext.asyncio import AsyncSession

from app.models.catalog import Product, ProductVariant, VariantPrice
from app.models.user import CustomerSegment, SegmentType, User
from tests.conftest import auth_header

URL = "/api/v1/admin/pricing/import-excel"
LABELS = {1: "HAM MALİYET", 2: "KDV DHL", 3: "ÖZEL FİYAT %18", 6: "BAYİ ", 9: "TOPTAN", 12: "KAFE", 15: "MAĞAZA", 17: "PERAKENDE"}


def _block(ws, row: int, base: int, name: str, prices: tuple, *, bayi=True, foreign=False, vat_excluded=False):
    """A four-row block as the workbook has it; `row` is 1-based, `base` the 0-based name column."""
    put = lambda r, c, v: ws.cell(r, base + c + 1, v)  # noqa: E731
    put(row, 0, name or None)
    for col, label in LABELS.items():
        if col != 6 or bayi:
            put(row, col, label)
    if foreign:
        put(row, 19, "TL")
    put(row + 1, 1, 100)
    put(row + 1, 2, "HARİÇ" if vat_excluded else 120)
    for col, price in zip((6, 9, 12, 17), prices):
        put(row + 1, col, price * 1.013)  # the unrounded chain
        put(row + 2, col, price)  # the rounded prices: two rows below the name


def _workbook(blocks) -> bytes:
    wb = openpyxl.Workbook()
    ws = wb.active
    ws.title = "hesaplama"
    ws.cell(1, 1, "SULTAN ÇANTALI")  # a section title, not a product
    for args, kwargs in blocks:
        _block(ws, *args, **kwargs)
    out = io.BytesIO()
    wb.save(out)
    return out.getvalue()


@pytest.fixture
async def products(db_session: AsyncSession, segments: dict[SegmentType, CustomerSegment]):
    tag = uuid.uuid4().hex[:6].upper()
    made = []

    async def product(code: str, name: str, skus: list[str]):
        p = Product(base_code=code, name=name, tags=[])
        db_session.add(p)
        await db_session.flush()
        variants = [ProductVariant(product_id=p.id, sku=sku, color_name=f"R{i}" if len(skus) > 1 else None, position=i) for i, sku in enumerate(skus)]
        db_session.add_all(variants)
        await db_session.flush()
        made.append(p)
        return p, variants

    hookah, (hookah_v,) = await product(f"LHT{tag}-V1", "Legend Test Masaüstü", [f"LHT{tag}-V1"])
    db_session.add(VariantPrice(
        variant_id=hookah_v.id, segment_id=segments[SegmentType.RETAIL].id, currency="TRY",
        price_excl_tax=Decimal("2500.00"), price_incl_tax=Decimal("3000.00"), vat_rate=Decimal("20.00"),
    ))
    # An old dollar row for the kafe group: it has to go once a TL price is written.
    db_session.add(VariantPrice(
        variant_id=hookah_v.id, segment_id=segments[SegmentType.CAFE].id, currency="USD",
        price_excl_tax=Decimal("50.00"), price_incl_tax=Decimal("60.00"), vat_rate=Decimal("20.00"),
    ))
    coloured, coloured_vs = await product(f"PBX{tag}", "Renkli Test Takımı", [f"SKU{tag} KR", f"SKU{tag} SY"])
    bottle, _ = await product(f"SSX{tag}", f"NTX{tag} Çantalı Nargile Şişesi", [f"SSX{tag}"])
    await db_session.commit()
    yield {"tag": tag, "hookah": hookah, "hookah_v": hookah_v, "coloured": coloured, "coloured_vs": coloured_vs, "bottle": bottle}

    ids = [p.id for p in made]
    variant_ids = select(ProductVariant.id).where(ProductVariant.product_id.in_(ids))
    await db_session.execute(delete(VariantPrice).where(VariantPrice.variant_id.in_(variant_ids)))
    await db_session.execute(delete(ProductVariant).where(ProductVariant.product_id.in_(ids)))
    await db_session.execute(delete(Product).where(Product.id.in_(ids)))
    await db_session.commit()


def _sheet(tag: str) -> bytes:
    return _workbook([
        ((2, 0, "", (100, 110, 120, 130)), {}),  # no name
        ((6, 0, f"LHT{tag} V1", (2155, 2545, 2930, 3450)), {}),  # code, space for the hyphen
        ((6, 19, f"sku{tag}kr", (1000, 1100, 1200, 1500)), {}),  # a colour's SKU, any case
        ((10, 0, f"NTX{tag}", (2100, 2475, 2850, 3350)), {}),  # only a bottle's *name* contains it
        ((10, 19, f"DUP{tag}", (1, 2, 3, 4)), {}),
        ((14, 0, f"DUP{tag}", (5, 6, 7, 8)), {}),  # the same name, other prices
        ((14, 19, f"EKS{tag}", (10, 20, 30, 0)), {}),  # no retail price
        ((18, 0, f"USD{tag}", (24, 25, 29, 38)), {"foreign": True}),
        ((22, 0, f"NOB{tag}", (0, 170, 180, 235)), {"bayi": False}),
        ((26, 0, f"HRC{tag}", (24, 25, 29, 38)), {"vat_excluded": True}),
    ])


async def _post(client: AsyncClient, admin: User, data: bytes, dry_run: bool):
    return await client.post(
        URL, params={"dry_run": str(dry_run).lower()}, headers=auth_header(admin),
        files={"file": ("01a-SULTAN.xlsx", data, "application/vnd.openxmlformats-officedocument.spreadsheetml.sheet")},
    )


async def _prices(db: AsyncSession, variant_id) -> dict:
    rows = (await db.execute(
        select(CustomerSegment.code, VariantPrice.currency, VariantPrice.price_incl_tax, VariantPrice.price_excl_tax)
        .join(CustomerSegment, CustomerSegment.id == VariantPrice.segment_id)
        .where(VariantPrice.variant_id == variant_id)
        .execution_options(populate_existing=True)
    )).all()
    return {(code, currency): (incl, excl) for code, currency, incl, excl in rows}


async def test_dry_run_reports_and_changes_nothing(client: AsyncClient, admin_user: User, products, db_session):
    tag = products["tag"]
    resp = await _post(client, admin_user, _sheet(tag), dry_run=True)
    assert resp.status_code == 200, resp.text
    body = resp.json()
    assert body["dry_run"] is True and body["updated_products"] == 0

    matched = {m["name"]: m for m in body["matched"] if tag in m["name"].upper()}
    assert set(matched) == {f"LHT{tag} V1", f"sku{tag}kr"}
    hookah = matched[f"LHT{tag} V1"]
    assert hookah["product_id"] == str(products["hookah"].id) and hookah["matched_by"] == "code"
    assert hookah["section"] == "SULTAN ÇANTALI" and hookah["row"] == 6 and hookah["side"] == "sol"
    assert hookah["prices"] == {"DEALER": "2155.00", "WHOLESALE": "2545.00", "CAFE": "2930.00", "RETAIL": "3450.00"}
    assert hookah["old_prices"]["RETAIL"] == {"currency": "TRY", "price_incl_tax": "3000.00"}
    coloured = matched[f"sku{tag}kr"]
    assert coloured["matched_by"] == "sku" and coloured["variants"] == 2 and coloured["side"] == "sağ"

    reasons = {(u["name"], u["row"], u["side"]): u["reason"] for u in body["unmatched"]}
    assert reasons[("", 2, "sol")] == "unnamed"
    # Never the bottle whose name merely contains the code.
    assert reasons[(f"NTX{tag}", 10, "sol")] == "not_found"
    assert str(products["bottle"].id) not in {m["product_id"] for m in body["matched"]}
    assert reasons[(f"DUP{tag}", 10, "sağ")] == reasons[(f"DUP{tag}", 14, "sol")] == "conflict"
    assert reasons[(f"EKS{tag}", 14, "sağ")] == "incomplete"

    skipped = {s["name"]: s["reason"] for s in body["skipped"]}
    assert skipped[f"USD{tag}"] == "foreign_currency"
    assert skipped[f"NOB{tag}"] == "no_dealer_price"
    assert skipped[f"HRC{tag}"] == "vat_excluded"

    assert await _prices(db_session, products["hookah_v"].id) == {
        ("RETAIL", "TRY"): (Decimal("3000.00"), Decimal("2500.00")),
        ("CAFE", "USD"): (Decimal("60.00"), Decimal("50.00")),
    }


async def test_apply_writes_all_four_groups(client: AsyncClient, admin_user: User, products, db_session):
    tag = products["tag"]
    resp = await _post(client, admin_user, _sheet(tag), dry_run=False)
    assert resp.status_code == 200, resp.text
    body = resp.json()
    assert body["dry_run"] is False
    assert body["updated_products"] == len(body["matched"]) and body["updated_variants"] >= 3

    # TL, VAT included in the sheet; the other side from the product's own 20%.
    assert await _prices(db_session, products["hookah_v"].id) == {
        ("DEALER", "TRY"): (Decimal("2155.00"), Decimal("1795.83")),
        ("WHOLESALE", "TRY"): (Decimal("2545.00"), Decimal("2120.83")),
        ("CAFE", "TRY"): (Decimal("2930.00"), Decimal("2441.67")),  # the dollar row is gone
        ("RETAIL", "TRY"): (Decimal("3450.00"), Decimal("2875.00")),
    }
    # Matched by one colour's SKU: every colour of the product is priced.
    for variant in products["coloured_vs"]:
        prices = await _prices(db_session, variant.id)
        assert prices[("RETAIL", "TRY")] == (Decimal("1500.00"), Decimal("1250.00"))
        assert prices[("DEALER", "TRY")][0] == Decimal("1000.00")
    assert await _prices(db_session, (await db_session.scalar(
        select(ProductVariant.id).where(ProductVariant.product_id == products["bottle"].id)
    ))) == {}


async def test_only_admins_and_only_workbooks(client: AsyncClient, admin_user: User, customer_user: User):
    data = _workbook([])
    assert (await client.post(URL, files={"file": ("a.xlsx", data)})).status_code == 401
    assert (await client.post(URL, headers=auth_header(customer_user), files={"file": ("a.xlsx", data)})).status_code == 403

    not_zip = await client.post(URL, headers=auth_header(admin_user), files={"file": ("a.xlsx", b"just text")})
    assert not_zip.status_code == 422
    archive = io.BytesIO()
    with zipfile.ZipFile(archive, "w") as z:
        z.writestr("hello.txt", "no workbook here")
    not_xlsx = await client.post(URL, headers=auth_header(admin_user), files={"file": ("a.xlsx", archive.getvalue())})
    assert not_xlsx.status_code == 422 and "xlsx" in not_xlsx.json()["detail"]
